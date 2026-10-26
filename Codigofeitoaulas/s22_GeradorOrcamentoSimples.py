import ast
import re
import hashlib
import sqlite3
import tkinter as tk
import random
import secrets
import string
import time
import calendar
from datetime import datetime, date
from tkinter import ttk, messagebox, scrolledtext, filedialog


class AplicacaoGerada(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('Novo Projeto')
        self.geometry('1000x730')
        self.minsize(820, 560)
        self.cor_fundo = '#eef2f7'
        self.cor_cartao = '#ffffff'
        self.cor_texto = '#203044'
        self.cor_destaque = '#2d6cdf'
        self.configure(bg=self.cor_fundo)
        self.historico_telas = []
        self.usuario_logado = None
        self.nome_usuario_logado = ''
        self.nome_tela_login = 'Tela Principal'
        self._metadados_tabelas = {'Tela Principal': {'tabela': 'registros', 'pk': 'id', 'colunas': ['item_1', 'qtd_1', 'preco_1', 'item_2', 'qtd_2', 'preco_2', 'item_3', 'qtd_3', 'preco_3', 'item_4', 'qtd_4', 'preco_4', 'total'], 'titulos': {'item_1': 'Item 1', 'qtd_1': 'Qtd 1', 'preco_1': 'Preco 1', 'item_2': 'Item 2', 'qtd_2': 'Qtd 2', 'preco_2': 'Preco 2', 'item_3': 'Item 3', 'qtd_3': 'Qtd 3', 'preco_3': 'Preco 3', 'item_4': 'Item 4', 'qtd_4': 'Qtd 4', 'preco_4': 'Preco 4', 'total': 'Total'}}}
        self._configurar_estilos()
        self.conexao = sqlite3.connect('dados_sistema.db')

        self.notebook_principal = ttk.Notebook(self)
        self.aba_sistema = ttk.Frame(
            self.notebook_principal,
            style='App.TFrame',
        )
        self.aba_dados = ttk.Frame(
            self.notebook_principal,
            style='App.TFrame',
        )
        self.notebook_principal.add(
            self.aba_sistema, text='Programa'
        )
        self.notebook_principal.add(
            self.aba_dados, text='Dados cadastrados'
        )
        self.notebook_principal.pack(fill='both', expand=True)

        self.container = ttk.Frame(
            self.aba_sistema, style='App.TFrame'
        )
        self.container.pack(fill='both', expand=True)
        self.telas = {}
        self.telas['Tela Principal'] = Tela1_TelaPrincipal(self.container, self)
        self.telas['Tela Principal'].place(x=0, y=0, relwidth=1, relheight=1)
        self._criar_visualizador_dados()
        self.notebook_principal.bind(
            '<<NotebookTabChanged>>', self._ao_mudar_aba
        )
        self.mostrar_tela('Tela Principal', registrar=False)

    def _configurar_estilos(self):
        estilo = ttk.Style(self)
        try:
            estilo.theme_use('clam')
        except tk.TclError:
            pass
        estilo.configure(
            'App.TFrame', background=self.cor_fundo
        )
        estilo.configure(
            'TFrame', background=self.cor_fundo
        )
        estilo.configure(
            'TLabel', background=self.cor_fundo,
            foreground=self.cor_texto, font=('Segoe UI', 10)
        )
        estilo.configure(
            'Title.TLabel', background=self.cor_fundo,
            foreground=self.cor_texto,
            font=('Segoe UI Semibold', 19)
        )
        estilo.configure(
            'TButton', font=('Segoe UI Semibold', 10),
            padding=(11, 7)
        )
        estilo.configure(
            'Treeview', font=('Segoe UI', 10), rowheight=28
        )
        estilo.configure(
            'Treeview.Heading',
            font=('Segoe UI Semibold', 10)
        )

    def _estilo_componente(self, nome_estilo, tipo, fonte, cor_texto='', cor_fundo=''):
        estilo = ttk.Style(self)
        opcoes = {}
        if fonte:
            opcoes['font'] = fonte
        if cor_texto:
            opcoes['foreground'] = cor_texto
        if cor_fundo:
            opcoes['background'] = cor_fundo
            if tipo in {'TEntry', 'TCombobox', 'TSpinbox', 'Treeview'}:
                opcoes['fieldbackground'] = cor_fundo
            if tipo == 'TScale':
                opcoes['troughcolor'] = cor_fundo
        if opcoes:
            estilo.configure(nome_estilo, **opcoes)

        if tipo == 'TCombobox':
            mapa = {}
            if cor_fundo:
                mapa['fieldbackground'] = [('readonly', cor_fundo)]
            if cor_texto:
                mapa['foreground'] = [('readonly', cor_texto)]
            if mapa:
                estilo.map(nome_estilo, **mapa)

        if tipo == 'TButton' and cor_fundo:
            estilo.map(nome_estilo, background=[('active', cor_fundo), ('pressed', cor_fundo)])

        if tipo == 'Treeview':
            cabecalho = nome_estilo + '.Heading'
            opcoes_cabecalho = {}
            if fonte:
                opcoes_cabecalho['font'] = fonte
            if cor_texto:
                opcoes_cabecalho['foreground'] = cor_texto
            if cor_fundo:
                opcoes_cabecalho['background'] = cor_fundo
            if opcoes_cabecalho:
                estilo.configure(cabecalho, **opcoes_cabecalho)

        if tipo == 'TLabelframe':
            rotulo = nome_estilo + '.Label'
            opcoes_rotulo = {}
            if fonte:
                opcoes_rotulo['font'] = fonte
            if cor_texto:
                opcoes_rotulo['foreground'] = cor_texto
            if cor_fundo:
                opcoes_rotulo['background'] = cor_fundo
            if opcoes_rotulo:
                estilo.configure(rotulo, **opcoes_rotulo)

        return nome_estilo

    def _criar_visualizador_dados(self):
        topo = ttk.Frame(self.aba_dados, padding=12)
        topo.pack(fill='x')
        ttk.Label(
            topo,
            text='Registros do sistema',
            style='Title.TLabel',
        ).pack(side='left')

        self.var_tabela_dados = tk.StringVar()
        self.combo_tabela_dados = ttk.Combobox(
            topo,
            textvariable=self.var_tabela_dados,
            state='readonly',
            width=34,
        )
        self.combo_tabela_dados.pack(
            side='left', padx=(18, 6)
        )
        self.combo_tabela_dados.configure(
            values=list(self._metadados_tabelas.keys())
        )
        self.combo_tabela_dados.bind(
            '<<ComboboxSelected>>',
            lambda e: self.atualizar_visualizador_dados(),
        )

        ttk.Button(
            topo, text='Atualizar',
            command=self.atualizar_visualizador_dados,
        ).pack(side='left', padx=5)
        ttk.Button(
            topo, text='Editar selecionado',
            command=self.editar_dado_selecionado,
        ).pack(side='left', padx=5)

        ttk.Button(
            topo, text='Excluir selecionado',
            command=self.excluir_dado_selecionado,
        ).pack(side='left', padx=5)

        barra_pesquisa = ttk.Frame(self.aba_dados, padding=(12, 0, 12, 10))
        barra_pesquisa.pack(fill='x')
        ttk.Label(barra_pesquisa, text='Pesquisar nos registros:').pack(side='left')
        self.var_pesquisa_dados = tk.StringVar()
        self.entry_pesquisa_dados = ttk.Entry(barra_pesquisa, textvariable=self.var_pesquisa_dados, width=42)
        self.entry_pesquisa_dados.pack(side='left', padx=(10, 6))
        self.var_pesquisa_dados.trace_add('write', lambda *_: self.atualizar_visualizador_dados())
        ttk.Button(
            barra_pesquisa, text='Limpar pesquisa',
            command=lambda: self.var_pesquisa_dados.set(''),
        ).pack(side='left', padx=5)
        ttk.Label(
            barra_pesquisa,
            text='Duplo clique em uma linha também abre a edição.',
        ).pack(side='right')

        corpo = ttk.Frame(
            self.aba_dados, padding=(12, 0, 12, 12)
        )
        corpo.pack(fill='both', expand=True)

        self.tree_dados = ttk.Treeview(
            corpo, show='headings'
        )
        scroll_y = ttk.Scrollbar(
            corpo, orient='vertical',
            command=self.tree_dados.yview,
        )
        scroll_x = ttk.Scrollbar(
            corpo, orient='horizontal',
            command=self.tree_dados.xview,
        )
        self.tree_dados.configure(
            yscrollcommand=scroll_y.set,
            xscrollcommand=scroll_x.set,
        )
        self.tree_dados.bind('<Double-1>', lambda e: self.editar_dado_selecionado())

        self.tree_dados.grid(
            row=0, column=0, sticky='nsew'
        )
        scroll_y.grid(row=0, column=1, sticky='ns')
        scroll_x.grid(row=1, column=0, sticky='ew')
        corpo.rowconfigure(0, weight=1)
        corpo.columnconfigure(0, weight=1)

        self.lbl_sem_tabelas = ttk.Label(
            corpo,
            text=(
                'Adicione campos ao projeto e use um botão '
                'com a ação Cadastrar no SQLite.'
            ),
        )

        if self._metadados_tabelas:
            primeira = next(iter(self._metadados_tabelas))
            self.var_tabela_dados.set(primeira)
            self.atualizar_visualizador_dados()
        else:
            self.lbl_sem_tabelas.place(
                relx=0.5, rely=0.5, anchor='center'
            )

    def _ao_mudar_aba(self, _evento=None):
        try:
            if self.notebook_principal.index('current') == 1:
                self.atualizar_visualizador_dados()
        except tk.TclError:
            pass

    def atualizar_visualizador_dados(self, tabela_preferida=None):
        if not self._metadados_tabelas:
            return

        if tabela_preferida:
            for rotulo, meta in self._metadados_tabelas.items():
                if meta['tabela'] == tabela_preferida:
                    self.var_tabela_dados.set(rotulo)
                    break

        rotulo = self.var_tabela_dados.get()
        meta = self._metadados_tabelas.get(rotulo)
        if not meta:
            return

        colunas = meta['colunas']
        consulta_colunas = [meta['pk']] + colunas
        self.tree_dados.configure(
            columns=colunas, show='headings'
        )

        for coluna in colunas:
            self.tree_dados.heading(
                coluna,
                text=meta.get('titulos', {}).get(coluna, coluna.replace('_', ' ').title()),
            )
            self.tree_dados.column(
                coluna,
                width=150 if coluna != meta['pk'] else 70,
                anchor='w',
            )

        for item in self.tree_dados.get_children():
            self.tree_dados.delete(item)

        sql = (
            f"SELECT {','.join(consulta_colunas)} "
            f"FROM {meta['tabela']} "
            f"ORDER BY {meta['pk']} DESC"
        )
        try:
            dados = self.conexao.execute(sql).fetchall()
        except sqlite3.Error:
            dados = []

        termo_pesquisa = getattr(self, 'var_pesquisa_dados', tk.StringVar()).get().strip().lower()
        if termo_pesquisa:
            dados = [
                linha for linha in dados
                if any(termo_pesquisa in str(valor).lower() for valor in linha[1:])
            ]

        for linha in dados:
            valores_visiveis = list(linha[1:])
            for indice_coluna, coluna in enumerate(colunas):
                if coluna.lower() == 'senha' and indice_coluna < len(valores_visiveis):
                    valores_visiveis[indice_coluna] = '••••••••'
            self.tree_dados.insert('', tk.END, iid=str(linha[0]), values=valores_visiveis)

    def editar_dado_selecionado(self):
        rotulo = self.var_tabela_dados.get()
        meta = self._metadados_tabelas.get(rotulo)
        selecao = self.tree_dados.selection()
        if not meta or not selecao:
            messagebox.showwarning('Editar', 'Selecione um registro na tabela.')
            return

        identificador = selecao[0]
        colunas = list(meta['colunas'])
        if not colunas:
            return
        sql = f"SELECT {','.join(colunas)} FROM {meta['tabela']} WHERE {meta['pk']} = ?"
        registro = self.conexao.execute(sql, (identificador,)).fetchone()
        if registro is None:
            messagebox.showwarning('Editar', 'O registro selecionado não foi encontrado.')
            self.atualizar_visualizador_dados()
            return

        janela = tk.Toplevel(self)
        janela.title('Editar registro')
        janela.transient(self)
        janela.grab_set()
        altura = min(700, max(470, 245 + min(len(colunas), 7) * 52))
        janela.geometry(f'760x{altura}')
        janela.minsize(640, 440)

        ttk.Label(janela, text='Editar registro', style='Title.TLabel').pack(anchor='w', padx=22, pady=(18, 4))
        ttk.Label(janela, text='Altere os campos desejados e clique em Salvar alterações.').pack(anchor='w', padx=22, pady=(0, 12))

        botoes = ttk.Frame(janela, padding=(22, 12, 22, 18))
        botoes.pack(side='bottom', fill='x')
        ttk.Separator(janela, orient='horizontal').pack(side='bottom', fill='x')

        area = ttk.Frame(janela)
        area.pack(fill='both', expand=True, padx=22, pady=(0, 8))
        canvas = tk.Canvas(area, highlightthickness=0)
        barra = ttk.Scrollbar(area, orient='vertical', command=canvas.yview)
        formulario = ttk.Frame(canvas)
        formulario.bind('<Configure>', lambda e: canvas.configure(scrollregion=canvas.bbox('all'))) 
        janela_canvas = canvas.create_window((0, 0), window=formulario, anchor='nw')
        canvas.bind('<Configure>', lambda e: canvas.itemconfigure(janela_canvas, width=e.width))
        canvas.configure(yscrollcommand=barra.set)
        canvas.pack(side='left', fill='both', expand=True)
        barra.pack(side='right', fill='y')
        canvas.bind('<Enter>', lambda e: canvas.bind_all('<MouseWheel>', lambda ev: canvas.yview_scroll(int(-1 * (ev.delta / 120)), 'units')))
        canvas.bind('<Leave>', lambda e: canvas.unbind_all('<MouseWheel>'))

        entradas = {}
        titulos = meta.get('titulos', {})
        for indice, (coluna, valor) in enumerate(zip(colunas, registro)):
            titulo = titulos.get(coluna, coluna.replace('_', ' ').title())
            ttk.Label(formulario, text=titulo).grid(row=indice, column=0, sticky='w', padx=(0, 14), pady=7)
            var = tk.StringVar()
            if coluna.lower() != 'senha':
                var.set('' if valor is None else str(valor))
            entrada = ttk.Entry(formulario, textvariable=var, show='•' if coluna.lower() == 'senha' else '')
            entrada.grid(row=indice, column=1, sticky='ew', pady=7)
            entradas[coluna] = var
            if coluna.lower() == 'senha':
                ttk.Label(formulario, text='Deixe vazio para manter a senha atual.').grid(row=indice, column=2, sticky='w', padx=(10, 0), pady=7)
        formulario.columnconfigure(1, weight=1)

        def salvar_edicao():
            sets = []
            valores = []
            for coluna in colunas:
                valor = entradas[coluna].get()
                if coluna.lower() == 'senha':
                    if not valor:
                        continue
                    valor = hashlib.sha256(valor.encode('utf-8')).hexdigest()
                if coluna.lower() == 'usuario' and not valor.strip():
                    messagebox.showwarning('Editar', 'O campo Usuário não pode ficar vazio.', parent=janela)
                    return
                sets.append(f'{coluna} = ?')
                valores.append(valor)
            if not sets:
                janela.destroy()
                return
            if meta['tabela'] == 'usuarios' and 'usuario' in colunas:
                usuario_novo = entradas['usuario'].get().strip()
                duplicado = self.conexao.execute('SELECT 1 FROM usuarios WHERE usuario = ? AND id <> ? LIMIT 1', (usuario_novo, identificador)).fetchone()
                if duplicado:
                    messagebox.showwarning('Editar', 'Esse nome de usuário já está sendo utilizado.', parent=janela)
                    return
            try:
                self.conexao.execute(
                    f"UPDATE {meta['tabela']} SET {', '.join(sets)} WHERE {meta['pk']} = ?",
                    valores + [identificador],
                )
                self.conexao.commit()
            except sqlite3.Error as erro:
                messagebox.showerror('Editar', f'Não foi possível salvar as alterações: {erro}', parent=janela)
                return
            janela.destroy()
            self.atualizar_visualizador_dados()
            for tela in self.telas.values():
                try:
                    tela.carregar_dados()
                except Exception:
                    pass
            messagebox.showinfo('Editar', 'Registro atualizado com sucesso!')

        ttk.Button(botoes, text='Salvar alterações', command=salvar_edicao, style='Accent.TButton').pack(side='right', padx=(8, 0))
        ttk.Button(botoes, text='Cancelar', command=janela.destroy).pack(side='right')

    def excluir_dado_selecionado(self):
        rotulo = self.var_tabela_dados.get()
        meta = self._metadados_tabelas.get(rotulo)
        selecao = self.tree_dados.selection()

        if not meta or not selecao:
            messagebox.showwarning(
                'Excluir',
                'Selecione um registro na tabela.',
            )
            return

        valores = self.tree_dados.item(
            selecao[0], 'values'
        )
        if not valores:
            return

        if not messagebox.askyesno(
            'Excluir',
            'Deseja excluir definitivamente o item selecionado do SQLite?',
        ):
            return

        self.conexao.execute(
            f"DELETE FROM {meta['tabela']} "
            f"WHERE {meta['pk']} = ?",
            (selecao[0],),
        )
        self.conexao.commit()
        self.atualizar_visualizador_dados()

        for tela in self.telas.values():
            try:
                tela.carregar_dados()
            except Exception:
                pass

    def mostrar_tela(self, nome, registrar=True):
        tela = self.telas.get(nome)
        if tela is None:
            messagebox.showwarning(
                'Navegação',
                f'A tela {nome} não existe neste projeto.',
            )
            return

        atual = getattr(self, '_tela_atual', None)
        if registrar and atual and atual != nome:
            self.historico_telas.append(atual)

        self._tela_atual = nome
        self.notebook_principal.select(self.aba_sistema)
        tela.lift()
        try:
            tela.carregar_dados()
        except Exception:
            pass

    def voltar_tela(self):
        if self.historico_telas:
            nome = self.historico_telas.pop()
            self.mostrar_tela(nome, registrar=False)

    def sair_da_conta(self):
        self.usuario_logado = None
        self.nome_usuario_logado = ''
        self.historico_telas.clear()
        self.mostrar_tela(
            self.nome_tela_login, registrar=False
        )

    def fechar(self):
        try:
            self.conexao.close()
        finally:
            self.destroy()


class Tela1_TelaPrincipal(ttk.Frame):
    def __init__(self, master, app):
        super().__init__(master, style='App.TFrame')
        self.app = app
        self._id_selecionado = None
        self._criar_banco()
        self._criar_interface()
        self._atualizar_contadores()


    def _criar_area_rolavel(self, largura, altura):
        self._largura_area = largura
        self._altura_area = altura
        self.canvas_rolavel = tk.Canvas(self, highlightthickness=0, bg=self.app.cor_fundo)
        self.barra_vertical = ttk.Scrollbar(self, orient='vertical', command=self.canvas_rolavel.yview)
        self.canvas_rolavel.configure(yscrollcommand=self.barra_vertical.set)
        self.canvas_rolavel.pack(side='left', fill='both', expand=True)
        self.barra_vertical.pack(side='right', fill='y')
        self.area = ttk.Frame(self.canvas_rolavel, style='App.TFrame', width=largura, height=altura)
        self.area.pack_propagate(False)
        self._janela_area = self.canvas_rolavel.create_window((0, 8), window=self.area, anchor='nw')
        self.canvas_rolavel.bind('<Configure>', self._ajustar_area_rolavel)
        self.area.bind('<Configure>', self._atualizar_scroll_area)
        self.canvas_rolavel.bind('<Enter>', lambda e: self.canvas_rolavel.bind_all('<MouseWheel>', self._ao_rolar_mouse))
        self.canvas_rolavel.bind('<Leave>', lambda e: self.canvas_rolavel.unbind_all('<MouseWheel>'))
        self.after(60, self._ajustar_area_rolavel)

    def _ajustar_area_rolavel(self, event=None):
        largura_canvas = max(1, self.canvas_rolavel.winfo_width())
        x = 0 if largura_canvas <= self._largura_area else int((largura_canvas - self._largura_area) / 2)
        self.canvas_rolavel.coords(self._janela_area, x, 8)
        self._atualizar_scroll_area()

    def _atualizar_scroll_area(self, event=None):
        largura_total = max(self._largura_area, self.canvas_rolavel.winfo_width())
        altura_total = max(self._altura_area + 16, self.canvas_rolavel.winfo_height())
        self.canvas_rolavel.configure(scrollregion=(0, 0, largura_total, altura_total))

    def _ao_rolar_mouse(self, event):
        try:
            deslocamento = int(-1 * (event.delta / 120))
            if deslocamento != 0:
                self.canvas_rolavel.yview_scroll(deslocamento, 'units')
        except Exception:
            pass

    def _criar_banco(self):
        self.app.conexao.execute('CREATE TABLE IF NOT EXISTS registros (id INTEGER PRIMARY KEY AUTOINCREMENT, item_1 TEXT, qtd_1 TEXT, preco_1 TEXT, item_2 TEXT, qtd_2 TEXT, preco_2 TEXT, item_3 TEXT, qtd_3 TEXT, preco_3 TEXT, item_4 TEXT, qtd_4 TEXT, preco_4 TEXT, total TEXT)')
        colunas_existentes = {linha[1] for linha in self.app.conexao.execute('PRAGMA table_info(registros)').fetchall()}
        if 'item_1' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN item_1 TEXT')
        if 'qtd_1' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN qtd_1 TEXT')
        if 'preco_1' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN preco_1 TEXT')
        if 'item_2' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN item_2 TEXT')
        if 'qtd_2' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN qtd_2 TEXT')
        if 'preco_2' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN preco_2 TEXT')
        if 'item_3' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN item_3 TEXT')
        if 'qtd_3' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN qtd_3 TEXT')
        if 'preco_3' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN preco_3 TEXT')
        if 'item_4' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN item_4 TEXT')
        if 'qtd_4' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN qtd_4 TEXT')
        if 'preco_4' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN preco_4 TEXT')
        if 'total' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN total TEXT')
        self.app.conexao.commit()

    def _criar_interface(self):
        self._criar_area_rolavel(1000, 650)
        self.label_1 = ttk.Label(self.area, text='Gerador de Orçamento Simples', style=self.app._estilo_componente('Componenteitem_1128213546.TLabel', 'TLabel', ('Segoe UI', 18), '', ''))
        self.label_1.place(x=38, y=40, width=417, height=30)
        self.label_2 = ttk.Label(self.area, text='Item 1', style='TLabel')
        self.label_2.place(x=38, y=108, width=180, height=30)
        self.var_item_1 = tk.StringVar()
        self.item_1 = ttk.Entry(self.area, textvariable=self.var_item_1)
        self.item_1.place(x=108, y=104, width=220, height=34)
        self.label_3 = ttk.Label(self.area, text='Qtd.', style='TLabel')
        self.label_3.place(x=365, y=104, width=180, height=30)
        self.var_qtd_1 = tk.StringVar()
        self.qtd_1 = ttk.Entry(self.area, textvariable=self.var_qtd_1)
        self.qtd_1.place(x=413, y=104, width=220, height=34)
        self.label_4 = ttk.Label(self.area, text='Preço', style='TLabel')
        self.label_4.place(x=668, y=104, width=180, height=30)
        self.var_preco_1 = tk.StringVar()
        self.preco_1 = ttk.Entry(self.area, textvariable=self.var_preco_1)
        self.preco_1.place(x=722, y=100, width=220, height=34)
        self.label_2_copia = ttk.Label(self.area, text='Item 2', style='TLabel')
        self.label_2_copia.place(x=38, y=172, width=180, height=30)
        self.var_item_2 = tk.StringVar()
        self.item_2 = ttk.Entry(self.area, textvariable=self.var_item_2)
        self.item_2.place(x=108, y=168, width=220, height=34)
        self.label_3_copia = ttk.Label(self.area, text='Qtd.', style='TLabel')
        self.label_3_copia.place(x=365, y=168, width=180, height=30)
        self.var_qtd_2 = tk.StringVar()
        self.qtd_2 = ttk.Entry(self.area, textvariable=self.var_qtd_2)
        self.qtd_2.place(x=413, y=168, width=220, height=34)
        self.label_4_copia = ttk.Label(self.area, text='Preço', style='TLabel')
        self.label_4_copia.place(x=668, y=168, width=180, height=30)
        self.var_preco_2 = tk.StringVar()
        self.preco_2 = ttk.Entry(self.area, textvariable=self.var_preco_2)
        self.preco_2.place(x=722, y=164, width=220, height=34)
        self.label_2_copia_copia = ttk.Label(self.area, text='Item 3', style='TLabel')
        self.label_2_copia_copia.place(x=38, y=230, width=180, height=30)
        self.var_item_3 = tk.StringVar()
        self.item_3 = ttk.Entry(self.area, textvariable=self.var_item_3)
        self.item_3.place(x=108, y=226, width=220, height=34)
        self.label_3_copia_copia = ttk.Label(self.area, text='Qtd.', style='TLabel')
        self.label_3_copia_copia.place(x=365, y=226, width=180, height=30)
        self.var_qtd_3 = tk.StringVar()
        self.qtd_3 = ttk.Entry(self.area, textvariable=self.var_qtd_3)
        self.qtd_3.place(x=413, y=226, width=220, height=34)
        self.label_4_copia_copia = ttk.Label(self.area, text='Preço', style='TLabel')
        self.label_4_copia_copia.place(x=668, y=226, width=180, height=30)
        self.var_preco_3 = tk.StringVar()
        self.preco_3 = ttk.Entry(self.area, textvariable=self.var_preco_3)
        self.preco_3.place(x=722, y=222, width=220, height=34)
        self.label_2_copia_copia_copia = ttk.Label(self.area, text='Item 4', style='TLabel')
        self.label_2_copia_copia_copia.place(x=38, y=281, width=180, height=30)
        self.var_item_4 = tk.StringVar()
        self.item_4 = ttk.Entry(self.area, textvariable=self.var_item_4)
        self.item_4.place(x=108, y=277, width=220, height=34)
        self.label_3_copia_copia_copia = ttk.Label(self.area, text='Qtd.', style='TLabel')
        self.label_3_copia_copia_copia.place(x=365, y=277, width=180, height=30)
        self.var_qtd_4 = tk.StringVar()
        self.qtd_4 = ttk.Entry(self.area, textvariable=self.var_qtd_4)
        self.qtd_4.place(x=413, y=277, width=220, height=34)
        self.label_4_copia_copia_copia = ttk.Label(self.area, text='Preço', style='TLabel')
        self.label_4_copia_copia_copia.place(x=668, y=277, width=180, height=30)
        self.var_preco_4 = tk.StringVar()
        self.preco_4 = ttk.Entry(self.area, textvariable=self.var_preco_4)
        self.preco_4.place(x=722, y=273, width=220, height=34)
        self.label_14 = ttk.Label(self.area, text='Total', style='TLabel')
        self.label_14.place(x=351, y=346, width=180, height=30)
        self.var_total = tk.StringVar()
        self.total = ttk.Entry(self.area, textvariable=self.var_total)
        self.total.place(x=413, y=342, width=220, height=34)
        self.btn_calcular = ttk.Button(self.area, text='Calcular orçamento', command=lambda: self.calcular_campos(['qtd_1', 'preco_1', 'qtd_2', 'preco_2', 'qtd_3', 'preco_3', 'qtd_4', 'preco_4'], 'total', 'Expressão personalizada', 'A*B + C*D + E*F + G*H', '', '', ''))
        self.btn_calcular.place(x=413, y=391, width=218, height=38)
        self.btn_salvar = ttk.Button(self.area, text='Salvar no SQLite', command=lambda: self.salvar_registro(['item_1', 'preco_1', 'item_2', 'preco_2', 'item_3', 'preco_3', 'item_4', 'preco_4', 'total']))
        self.btn_salvar.place(x=38, y=454, width=288, height=38)

    def _valor_componente(self, nome):
        variavel = getattr(self, f'var_{nome}', None)
        if variavel is not None:
            return variavel.get()
        widget = getattr(self, nome, None)
        if isinstance(widget, tk.Text):
            return widget.get('1.0', 'end-1c')
        if isinstance(widget, tk.Listbox):
            selecao = widget.curselection()
            if selecao:
                return widget.get(selecao[0])
            return '\n'.join(widget.get(0, tk.END))
        return ''

    def _definir_valor_componente(self, nome, valor):
        valor = '' if valor is None else valor
        variavel = getattr(self, f'var_{nome}', None)
        if variavel is not None:
            variavel.set(valor)
            return True
        widget = getattr(self, nome, None)
        if isinstance(widget, tk.Text):
            widget.delete('1.0', tk.END)
            texto = '\n'.join(str(x) for x in valor) if isinstance(valor, (list, tuple)) else str(valor)
            widget.insert('1.0', texto)
            return True
        if isinstance(widget, tk.Listbox):
            widget.delete(0, tk.END)
            itens = valor if isinstance(valor, (list, tuple)) else str(valor).splitlines()
            for item in itens:
                widget.insert(tk.END, item)
            return True
        return False

    def _atualizar_contadores(self):
        pass

    def exportar_dados_excel(self, alvos=None):
        try:
            from openpyxl import Workbook
        except ImportError:
            messagebox.showerror('Excel', 'Para exportar arquivos .xlsx, instale a biblioteca openpyxl com: pip install openpyxl')
            return
        nomes = list(alvos or [])
        if not nomes:
            nomes = [nome for nome, widget in self.__dict__.items() if isinstance(widget, (ttk.Treeview, tk.Listbox))]
        fontes = [(nome, getattr(self, nome, None)) for nome in nomes]
        fontes = [(nome, widget) for nome, widget in fontes if isinstance(widget, (ttk.Treeview, tk.Listbox))]
        if not fontes:
            messagebox.showwarning('Excel', 'Selecione uma Treeview ou uma Listbox nos Campos da ação.')
            return
        caminho = filedialog.asksaveasfilename(title='Exportar para Excel', defaultextension='.xlsx', filetypes=[('Arquivo do Excel', '*.xlsx')])
        if not caminho:
            return
        wb = Workbook()
        wb.remove(wb.active)
        def valor_seguro(valor):
            if isinstance(valor, str) and valor[:1] in ('=', '+', '-', '@'):
                return "'" + valor
            return valor
        for indice, (nome, widget) in enumerate(fontes, 1):
            titulo = re.sub(r'[:\\/?*\[\]]', '_', str(nome))[:31] or f'Dados{indice}'
            base = titulo
            contador = 2
            while titulo in wb.sheetnames:
                sufixo = f'_{contador}'
                titulo = base[:31-len(sufixo)] + sufixo
                contador += 1
            ws = wb.create_sheet(titulo)
            if isinstance(widget, ttk.Treeview):
                colunas = list(widget['columns'])
                cabecalhos = [widget.heading(coluna, 'text') or str(coluna) for coluna in colunas]
                ws.append(cabecalhos)
                for item in widget.get_children():
                    ws.append([valor_seguro(v) for v in widget.item(item, 'values')])
            else:
                ws.append(['Item'])
                for item in widget.get(0, tk.END):
                    ws.append([valor_seguro(item)])
            for coluna in ws.columns:
                letra = coluna[0].column_letter
                maior = max((len(str(celula.value or '')) for celula in coluna), default=10)
                ws.column_dimensions[letra].width = min(max(maior + 2, 12), 60)
        try:
            wb.save(caminho)
        except OSError as erro:
            messagebox.showerror('Excel', f'Não foi possível salvar o arquivo: {erro}')
            return
        messagebox.showinfo('Excel', 'Dados exportados para o Excel com sucesso!')

    def _formatar_tempo(self, segundos):
        segundos = max(0.0, float(segundos))
        horas = int(segundos // 3600)
        minutos = int((segundos % 3600) // 60)
        resto = segundos % 60
        return f'{horas:02d}:{minutos:02d}:{resto:04.1f}'

    def cronometro_iniciar(self, destino):
        if getattr(self, '_cronometro_rodando', False):
            return
        self._cronometro_rodando = True
        self._cronometro_acumulado = float(getattr(self, '_cronometro_acumulado', 0.0))
        self._cronometro_inicio = time.perf_counter()
        self._cronometro_tick(destino)

    def _cronometro_tick(self, destino):
        if not getattr(self, '_cronometro_rodando', False):
            return
        atual = self._cronometro_acumulado + (time.perf_counter() - self._cronometro_inicio)
        self._definir_valor_componente(destino, self._formatar_tempo(atual))
        self.after(100, lambda: self._cronometro_tick(destino))

    def cronometro_pausar(self, destino):
        if getattr(self, '_cronometro_rodando', False):
            self._cronometro_acumulado += time.perf_counter() - self._cronometro_inicio
        self._cronometro_rodando = False
        self._definir_valor_componente(destino, self._formatar_tempo(getattr(self, '_cronometro_acumulado', 0.0)))

    def cronometro_zerar(self, destino):
        self._cronometro_rodando = False
        self._cronometro_acumulado = 0.0
        self._definir_valor_componente(destino, '00:00:00.0')

    def temporizador_iniciar(self, origem, destino):
        if getattr(self, '_temporizador_rodando', False):
            return
        try:
            valor = int(float(str(self._valor_componente(origem)).replace(',', '.')))
        except (TypeError, ValueError):
            valor = 0
        if valor <= 0:
            messagebox.showwarning('Temporizador', 'Informe uma quantidade de segundos maior que zero.')
            return
        self._temporizador_restante = valor
        self._temporizador_rodando = True
        self._temporizador_tick(destino)

    def _temporizador_tick(self, destino):
        if not getattr(self, '_temporizador_rodando', False):
            return
        restante = int(getattr(self, '_temporizador_restante', 0))
        self._definir_valor_componente(destino, self._formatar_tempo(restante))
        if restante <= 0:
            self._temporizador_rodando = False
            messagebox.showinfo('Temporizador', 'Tempo encerrado!')
            return
        self._temporizador_restante = restante - 1
        self.after(1000, lambda: self._temporizador_tick(destino))

    def temporizador_pausar(self, destino):
        self._temporizador_rodando = False
        self._definir_valor_componente(destino, self._formatar_tempo(getattr(self, '_temporizador_restante', 0)))

    def temporizador_zerar(self, destino):
        self._temporizador_rodando = False
        self._temporizador_restante = 0
        self._definir_valor_componente(destino, '00:00:00.0')

    def gerar_senha_segura(self, campo_tamanho, campo_simbolos, destino):
        try:
            tamanho = int(float(str(self._valor_componente(campo_tamanho)).replace(',', '.')))
        except (TypeError, ValueError):
            tamanho = 12
        tamanho = min(max(tamanho, 8), 128)
        usar_simbolos = bool(self._valor_componente(campo_simbolos)) if campo_simbolos else True
        alfabeto = string.ascii_letters + string.digits
        if usar_simbolos:
            alfabeto += '!@#$%&*_-+=?'
        senha = ''.join(secrets.choice(alfabeto) for _ in range(tamanho))
        self._definir_valor_componente(destino, senha)

    def calcular_idade_completa(self, origem, destino):
        texto = str(self._valor_componente(origem)).strip()
        try:
            nascimento = datetime.strptime(texto, '%d/%m/%Y').date()
        except ValueError:
            messagebox.showwarning('Idade', 'Informe a data de nascimento no formato dd/mm/aaaa.')
            return
        hoje = date.today()
        if nascimento > hoje:
            messagebox.showwarning('Idade', 'A data de nascimento não pode estar no futuro.')
            return
        anos = hoje.year - nascimento.year
        meses = hoje.month - nascimento.month
        dias = hoje.day - nascimento.day
        if dias < 0:
            meses -= 1
            mes_anterior = hoje.month - 1 or 12
            ano_mes_anterior = hoje.year if hoje.month > 1 else hoje.year - 1
            dias += calendar.monthrange(ano_mes_anterior, mes_anterior)[1]
        if meses < 0:
            anos -= 1
            meses += 12
        self._definir_valor_componente(destino, f'{anos} anos, {meses} meses e {dias} dias')

    def sortear_nome(self, origem, destino):
        widget = getattr(self, origem, None)
        if isinstance(widget, tk.Listbox):
            nomes = [str(x).strip() for x in widget.get(0, tk.END) if str(x).strip()]
        else:
            nomes = [x.strip() for x in str(self._valor_componente(origem)).splitlines() if x.strip()]
        if not nomes:
            messagebox.showwarning('Sorteio', 'A lista de nomes está vazia.')
            return
        self._definir_valor_componente(destino, random.choice(nomes))

    def gerar_numeros_sorteio(self, origem, destino):
        try:
            quantidade = int(float(str(self._valor_componente(origem)).replace(',', '.')))
        except (TypeError, ValueError):
            quantidade = 6
        if quantidade <= 0:
            quantidade = 6
        quantidade = min(quantidade, 60)
        numeros = sorted(random.sample(range(1, 61), quantidade))
        self._definir_valor_componente(destino, [f'{numero:02d}' for numero in numeros])
        self._atualizar_contadores()

    def validar_formulario(self):
        return True

    def _aplicar_mascara(self, nome, mascara):
        valor = str(self._valor_componente(nome))
        digitos = re.sub(r'\D', '', valor)
        novo = valor
        if mascara == 'CPF' and len(digitos) >= 11:
            d = digitos[:11]; novo = f'{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}'
        elif mascara == 'CNPJ' and len(digitos) >= 14:
            d = digitos[:14]; novo = f'{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}'
        elif mascara == 'Telefone' and len(digitos) >= 10:
            d = digitos[:11]; novo = f'({d[:2]}) {d[2:7]}-{d[7:]}' if len(d) == 11 else f'({d[:2]}) {d[2:6]}-{d[6:]}'
        elif mascara == 'CEP' and len(digitos) >= 8:
            d = digitos[:8]; novo = f'{d[:5]}-{d[5:]}'
        elif mascara == 'Data' and len(digitos) >= 8:
            d = digitos[:8]; novo = f'{d[:2]}/{d[2:4]}/{d[4:]}'
        elif mascara == 'Hora' and len(digitos) >= 4:
            d = digitos[:4]; novo = f'{d[:2]}:{d[2:4]}'
        elif mascara == 'Moeda':
            try:
                numero = float(valor.replace('R$', '').replace('.', '').replace(',', '.').strip())
                novo = f'R$ {numero:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
            except ValueError:
                return
        self._definir_valor_componente(nome, novo)

    def validar_login(self):
        messagebox.showinfo('Login', 'Esta tela não está configurada como Login.')

    def salvar_registro(self, selecionados=None):
        if not self.validar_formulario():
            return
        mapa = {'item_1': 'item_1', 'qtd_1': 'qtd_1', 'preco_1': 'preco_1', 'item_2': 'item_2', 'qtd_2': 'qtd_2', 'preco_2': 'preco_2', 'item_3': 'item_3', 'qtd_3': 'qtd_3', 'preco_3': 'preco_3', 'item_4': 'item_4', 'qtd_4': 'qtd_4', 'preco_4': 'preco_4', 'total': 'total'}
        escolhidos = list(selecionados or ['item_1', 'qtd_1', 'preco_1', 'item_2', 'qtd_2', 'preco_2', 'item_3', 'qtd_3', 'preco_3', 'item_4', 'qtd_4', 'preco_4', 'total'])
        escolhidos = [nome for nome in escolhidos if nome in mapa]
        if not escolhidos:
            messagebox.showwarning('SQLite', 'Selecione pelo menos um campo para cadastrar.')
            return
        campos = [mapa[nome] for nome in escolhidos]
        valores = []
        for nome, campo in zip(escolhidos, campos):
            valor = self._valor_componente(nome)
            if campo == 'senha' and str(valor):
                valor = hashlib.sha256(str(valor).encode('utf-8')).hexdigest()
            valores.append(valor)
        placeholders = ','.join('?' for _ in campos)
        sql = 'INSERT INTO registros (' + ','.join(campos) + ') VALUES (' + placeholders + ')' 
        try:
            self.app.conexao.execute(sql, valores)
            self.app.conexao.commit()
        except sqlite3.IntegrityError as erro:
            messagebox.showerror('Banco de dados', f'Não foi possível cadastrar: {erro}')
            return
        self.app.atualizar_visualizador_dados('registros')
        messagebox.showinfo('SQLite', 'Dados cadastrados com sucesso! Abra a aba Dados cadastrados para visualizar.')

    def carregar_dados(self):
        pass
        self._atualizar_contadores()

    def aplicar_filtro(self):
        pass
        self._atualizar_contadores()

    def preencher_formulario(self, event=None):
        pass

    def atualizar_registro(self, selecionados=None):
        messagebox.showinfo('Configuração', 'Para atualizar pela tela, inclua uma Treeview e selecione o registro.')

    def excluir_registro(self):
        messagebox.showinfo('Excluir', 'Abra a aba Dados cadastrados para editar ou excluir o registro selecionado.')

    def limpar_campos(self, selecionados=None):
        self._id_selecionado = None
        nomes = list(selecionados or ['item_1', 'qtd_1', 'preco_1', 'item_2', 'qtd_2', 'preco_2', 'item_3', 'qtd_3', 'preco_3', 'item_4', 'qtd_4', 'preco_4', 'total'])
        for nome in nomes:
            variavel = getattr(self, f'var_{nome}', None)
            if variavel is not None:
                try:
                    variavel.set(False if isinstance(variavel, tk.BooleanVar) else '')
                except tk.TclError:
                    pass
                continue
            widget = getattr(self, nome, None)
            if isinstance(widget, tk.Text):
                widget.delete('1.0', tk.END)
            elif isinstance(widget, tk.Listbox):
                widget.delete(0, tk.END)

    def _valor_para_formula(self, valor):
        if isinstance(valor, (int, float, bool)):
            return valor
        original = '' if valor is None else str(valor).strip()
        texto = original.replace('R$', '').replace(' ', '')
        try:
            if ',' in texto:
                texto = texto.replace('.', '').replace(',', '.')
            numero = float(texto)
            return int(numero) if numero.is_integer() else numero
        except (ValueError, TypeError):
            return original

    def _formatar_resultado(self, valor):
        if isinstance(valor, float):
            if valor.is_integer():
                return str(int(valor))
            return f'{valor:.10f}'.rstrip('0').rstrip('.').replace('.', ',')
        return str(valor)

    def _avaliar_expressao_segura(self, expressao, valores=None, resultado_anterior=None, n=None):
        valores = list(valores or [])
        def nome_campo_formula(indice):
            numero = indice + 1
            partes = []
            while numero > 0:
                numero, resto = divmod(numero - 1, 26)
                partes.append(chr(65 + resto))
            return ''.join(reversed(partes))
        ambiente = {nome_campo_formula(i): valor for i, valor in enumerate(valores)}
        ambiente.update({f'CAMPO{i + 1}': valor for i, valor in enumerate(valores)})
        if resultado_anterior is not None: ambiente['R'] = resultado_anterior
        if n is not None: ambiente['N'] = n
        funcoes = {
            'round': round, 'abs': abs, 'min': min, 'max': max,
            'int': int, 'float': float, 'str': str,
            'upper': lambda x: str(x).upper(),
            'lower': lambda x: str(x).lower(), 'len': len
        }
        arvore = ast.parse(str(expressao), mode='eval')

        def visitar(no):
            if isinstance(no, ast.Expression):
                return visitar(no.body)
            if isinstance(no, ast.Constant):
                return no.value
            if isinstance(no, ast.Name):
                if no.id in ambiente:
                    return ambiente[no.id]
                raise ValueError(f'Nome não permitido: {no.id}')
            if isinstance(no, ast.BinOp):
                a, b = visitar(no.left), visitar(no.right)
                if isinstance(no.op, ast.Add): return a + b
                if isinstance(no.op, ast.Sub): return a - b
                if isinstance(no.op, ast.Mult): return a * b
                if isinstance(no.op, ast.Div): return a / b
                if isinstance(no.op, ast.FloorDiv): return a // b
                if isinstance(no.op, ast.Mod): return a % b
                if isinstance(no.op, ast.Pow): return a ** b
                raise ValueError('Operador não permitido')
            if isinstance(no, ast.UnaryOp):
                valor = visitar(no.operand)
                if isinstance(no.op, ast.USub): return -valor
                if isinstance(no.op, ast.UAdd): return +valor
                if isinstance(no.op, ast.Not): return not valor
                raise ValueError('Operador unário não permitido')
            if isinstance(no, ast.BoolOp):
                valores_bool = [visitar(x) for x in no.values]
                return all(valores_bool) if isinstance(no.op, ast.And) else any(valores_bool)
            if isinstance(no, ast.Compare):
                esquerda = visitar(no.left)
                for operador, comparador in zip(no.ops, no.comparators):
                    direita = visitar(comparador)
                    if isinstance(operador, ast.Eq): ok = esquerda == direita
                    elif isinstance(operador, ast.NotEq): ok = esquerda != direita
                    elif isinstance(operador, ast.Gt): ok = esquerda > direita
                    elif isinstance(operador, ast.GtE): ok = esquerda >= direita
                    elif isinstance(operador, ast.Lt): ok = esquerda < direita
                    elif isinstance(operador, ast.LtE): ok = esquerda <= direita
                    else: raise ValueError('Comparação não permitida')
                    if not ok: return False
                    esquerda = direita
                return True
            if isinstance(no, ast.IfExp):
                return visitar(no.body) if visitar(no.test) else visitar(no.orelse)
            if isinstance(no, ast.Call):
                if not isinstance(no.func, ast.Name) or no.func.id not in funcoes:
                    raise ValueError('Função não permitida')
                if no.keywords:
                    raise ValueError('Argumentos nomeados não são permitidos')
                return funcoes[no.func.id](*[visitar(x) for x in no.args])
            raise ValueError(f'Elemento não permitido: {type(no).__name__}')

        return visitar(arvore)

    def _operacao_basica(self, valores, operador, expressao=''):
        if operador == 'Expressão personalizada':
            if not str(expressao).strip():
                raise ValueError('Informe a expressão personalizada.')
            return self._avaliar_expressao_segura(expressao, valores)
        if operador == 'Operação definida no Campo 3':
            if len(valores) < 3:
                raise ValueError('Selecione Campo 1, Campo 2 e Campo 3.')
            op = str(valores[2]).strip().lower()
            mapa = {
                'somar': 'Somar', '+': 'Somar',
                'subtrair': 'Subtrair', '-': 'Subtrair',
                'multiplicar': 'Multiplicar', '*': 'Multiplicar', '×': 'Multiplicar',
                'dividir': 'Dividir', '/': 'Dividir', '÷': 'Dividir'
            }
            operador = mapa.get(op, '')
            valores = valores[:2]
        if not valores:
            raise ValueError('Selecione pelo menos um campo.')
        if operador == 'Somar':
            return sum(valores)
        if operador == 'Subtrair':
            resultado = valores[0]
            for valor in valores[1:]: resultado -= valor
            return resultado
        if operador == 'Multiplicar':
            resultado = 1
            for valor in valores: resultado *= valor
            return resultado
        if operador == 'Dividir':
            resultado = valores[0]
            for valor in valores[1:]: resultado /= valor
            return resultado
        if operador == 'Média':
            return sum(valores) / len(valores)
        if operador == 'Porcentagem':
            if len(valores) < 2: raise ValueError('Porcentagem precisa de dois campos.')
            return valores[0] * valores[1] / 100
        if operador == 'Resto da divisão':
            if len(valores) < 2: raise ValueError('Resto precisa de dois campos.')
            return valores[0] % valores[1]
        if operador == 'Potência':
            if len(valores) < 2: raise ValueError('Potência precisa de dois campos.')
            return valores[0] ** valores[1]
        if operador == 'Mínimo':
            return min(valores)
        if operador == 'Máximo':
            return max(valores)
        raise ValueError('Escolha uma operação matemática.')

    def inserir_valor(self, destino, valor, modo='Adicionar'):
        if not destino:
            messagebox.showwarning('Ação', 'Selecione o campo que receberá o valor.')
            return
        atual = str(self._valor_componente(destino))
        novo = str(valor) if modo == 'Substituir' else atual + str(valor)
        if not self._definir_valor_componente(destino, novo):
            messagebox.showwarning('Ação', 'O campo de destino não existe nesta tela.')

    def calcular_campos(self, campos, destino, operador, expressao='', formato='', destino_2='', expressao_2=''):
        if not campos or not destino:
            messagebox.showwarning('Cálculo', 'Selecione os campos da operação e o campo de resultado.')
            return
        try:
            valores = [self._valor_para_formula(self._valor_componente(nome)) for nome in campos]
            resultado = self._operacao_basica(valores, operador, expressao)
            texto = self._formatar_resultado(resultado)
            if formato:
                texto = str(formato).replace('{resultado}', texto)
            if not self._definir_valor_componente(destino, texto):
                raise ValueError('Campo de resultado não encontrado.')
            if destino_2 and expressao_2:
                resultado_2 = self._avaliar_expressao_segura(expressao_2, valores, resultado_anterior=resultado)
                self._definir_valor_componente(destino_2, self._formatar_resultado(resultado_2))
        except ZeroDivisionError:
            messagebox.showwarning('Cálculo', 'Não é possível dividir por zero.')
        except Exception as erro:
            messagebox.showwarning('Cálculo', f'Não foi possível executar a operação: {erro}')

    def calcular_expressao_campo(self, origem, destino, auxiliar=''):
        if not origem or not destino:
            messagebox.showwarning('Cálculo', 'Selecione o campo da expressão e o campo de resultado.')
            return
        original = str(self._valor_componente(origem)).strip()
        expressao = original.replace('×', '*').replace('÷', '/').replace(',', '.')
        if not expressao:
            return
        try:
            resultado = self._avaliar_expressao_segura(expressao)
            if not isinstance(resultado, (int, float)):
                raise ValueError('A expressão deve produzir um número.')
            texto = self._formatar_resultado(resultado)
            self._definir_valor_componente(destino, texto)
            if auxiliar:
                widget = getattr(self, auxiliar, None)
                if isinstance(widget, tk.Listbox):
                    widget.insert(tk.END, f'{original} = {texto}')
                elif isinstance(widget, tk.Text):
                    widget.insert(tk.END, f'{original} = {texto}\n')
        except ZeroDivisionError:
            messagebox.showwarning('Cálculo', 'Não é possível dividir por zero.')
        except Exception as erro:
            messagebox.showwarning('Cálculo', f'Expressão inválida: {erro}')

    def aplicar_condicao(self, campos, destino, comparador, valor_fixo='', verdadeiro='Sim', falso='Não'):
        if not campos or not destino or not comparador:
            messagebox.showwarning('Condição', 'Selecione Campo 1, comparador e campo de resultado.')
            return
        a = self._valor_para_formula(self._valor_componente(campos[0]))
        if len(campos) >= 2:
            b = self._valor_para_formula(self._valor_componente(campos[1]))
        else:
            b = self._valor_para_formula(valor_fixo)
        try:
            if comparador == '==': ok = a == b
            elif comparador == '!=': ok = a != b
            elif comparador == '>': ok = a > b
            elif comparador == '>=': ok = a >= b
            elif comparador == '<': ok = a < b
            elif comparador == '<=': ok = a <= b
            else: raise ValueError('Comparador inválido')
            self._definir_valor_componente(destino, verdadeiro if ok else falso)
        except Exception as erro:
            messagebox.showwarning('Condição', f'Não foi possível comparar os valores: {erro}')

    def gerar_lista_repeticao(self, campos, destino, operador, inicio='1', fim='10', modelo=''):
        if not campos or not destino:
            messagebox.showwarning('Repetição', 'Selecione Campo 1 e o campo de resultado.')
            return
        try:
            a = self._valor_para_formula(self._valor_componente(campos[0]))
            inicio_i, fim_i = int(inicio), int(fim)
            passo = 1 if fim_i >= inicio_i else -1
            linhas_saida = []
            for n in range(inicio_i, fim_i + passo, passo):
                resultado = self._operacao_basica([a, n], operador, '')
                texto_resultado = self._formatar_resultado(resultado)
                linha = modelo or '{A} {operador} {N} = {resultado}'
                linha = str(linha).replace('{A}', self._formatar_resultado(a)).replace('{N}', str(n))
                linha = linha.replace('{operador}', operador).replace('{resultado}', texto_resultado)
                linhas_saida.append(linha)
            self._definir_valor_componente(destino, linhas_saida)
        except Exception as erro:
            messagebox.showwarning('Repetição', f'Não foi possível gerar a lista: {erro}')



if __name__ == '__main__':
    app = AplicacaoGerada()
    app.protocol('WM_DELETE_WINDOW', app.fechar)
    app.mainloop()
