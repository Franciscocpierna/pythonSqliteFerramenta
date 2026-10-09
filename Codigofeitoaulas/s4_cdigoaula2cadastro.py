import ast
import re
import hashlib
import sqlite3
import tkinter as tk
import random
import secrets
import string
import time
import calendar
from datetime import datetime, date
from tkinter import ttk, messagebox, scrolledtext, filedialog


class AplicacaoGerada(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('Novo Projeto')
        self.geometry('1131x828')
        self.minsize(820, 560)
        self.cor_fundo = '#eef2f7'
        self.cor_cartao = '#ffffff'
        self.cor_texto = '#203044'
        self.cor_destaque = '#2d6cdf'
        self.configure(bg=self.cor_fundo)
        self.historico_telas = []
        self.usuario_logado = None
        self.nome_usuario_logado = ''
        self.nome_tela_login = 'Tela Principal'
        self._metadados_tabelas = {'Tela Principal': {'tabela': 'registros', 'pk': 'id', 'colunas': ['nome', 'curso', 'idade', 'e_mail', 'sexo', 'telefone'], 'titulos': {'nome': 'Nome', 'curso': 'Curso', 'idade': 'Idade', 'e_mail': 'E Mail', 'sexo': 'Sexo', 'telefone': 'Telefone'}}}
        self._configurar_estilos()
        self.conexao = sqlite3.connect('dados_sistema.db')

        self.notebook_principal = ttk.Notebook(self)
        self.aba_sistema = ttk.Frame(
            self.notebook_principal,
            style='App.TFrame',
        )
        self.notebook_principal.add(
            self.aba_sistema, text='Programa'
        )
        self.notebook_principal.pack(fill='both', expand=True)

        self.container = ttk.Frame(
            self.aba_sistema, style='App.TFrame'
        )
        self.container.pack(fill='both', expand=True)
        self.telas = {}
        self.telas['Tela Principal'] = Tela1_TelaPrincipal(self.container, self)
        self.telas['Tela Principal'].place(x=0, y=0, relwidth=1, relheight=1)
        self.mostrar_tela('Tela Principal', registrar=False)

    def _configurar_estilos(self):
        estilo = ttk.Style(self)
        try:
            estilo.theme_use('clam')
        except tk.TclError:
            pass
        estilo.configure(
            'App.TFrame', background=self.cor_fundo
        )
        estilo.configure(
            'TFrame', background=self.cor_fundo
        )
        estilo.configure(
            'TLabel', background=self.cor_fundo,
            foreground=self.cor_texto, font=('Segoe UI', 10)
        )
        estilo.configure(
            'Title.TLabel', background=self.cor_fundo,
            foreground=self.cor_texto,
            font=('Segoe UI Semibold', 19)
        )
        estilo.configure(
            'TButton', font=('Segoe UI Semibold', 10),
            padding=(11, 7)
        )
        estilo.configure(
            'Treeview', font=('Segoe UI', 10), rowheight=28
        )
        estilo.configure(
            'Treeview.Heading',
            font=('Segoe UI Semibold', 10)
        )

    def _estilo_componente(self, nome_estilo, tipo, fonte, cor_texto='', cor_fundo=''):
        estilo = ttk.Style(self)
        opcoes = {}
        if fonte:
            opcoes['font'] = fonte
        if cor_texto:
            opcoes['foreground'] = cor_texto
        if cor_fundo:
            opcoes['background'] = cor_fundo
            if tipo in {'TEntry', 'TCombobox', 'TSpinbox', 'Treeview'}:
                opcoes['fieldbackground'] = cor_fundo
            if tipo == 'TScale':
                opcoes['troughcolor'] = cor_fundo
        if opcoes:
            estilo.configure(nome_estilo, **opcoes)

        if tipo == 'TCombobox':
            mapa = {}
            if cor_fundo:
                mapa['fieldbackground'] = [('readonly', cor_fundo)]
            if cor_texto:
                mapa['foreground'] = [('readonly', cor_texto)]
            if mapa:
                estilo.map(nome_estilo, **mapa)

        if tipo == 'TButton' and cor_fundo:
            estilo.map(nome_estilo, background=[('active', cor_fundo), ('pressed', cor_fundo)])

        if tipo == 'Treeview':
            cabecalho = nome_estilo + '.Heading'
            opcoes_cabecalho = {}
            if fonte:
                opcoes_cabecalho['font'] = fonte
            if cor_texto:
                opcoes_cabecalho['foreground'] = cor_texto
            if cor_fundo:
                opcoes_cabecalho['background'] = cor_fundo
            if opcoes_cabecalho:
                estilo.configure(cabecalho, **opcoes_cabecalho)

        if tipo == 'TLabelframe':
            rotulo = nome_estilo + '.Label'
            opcoes_rotulo = {}
            if fonte:
                opcoes_rotulo['font'] = fonte
            if cor_texto:
                opcoes_rotulo['foreground'] = cor_texto
            if cor_fundo:
                opcoes_rotulo['background'] = cor_fundo
            if opcoes_rotulo:
                estilo.configure(rotulo, **opcoes_rotulo)

        return nome_estilo

    def mostrar_tela(self, nome, registrar=True):
        tela = self.telas.get(nome)
        if tela is None:
            messagebox.showwarning(
                'Navegação',
                f'A tela {nome} não existe neste projeto.',
            )
            return

        atual = getattr(self, '_tela_atual', None)
        if registrar and atual and atual != nome:
            self.historico_telas.append(atual)

        self._tela_atual = nome
        self.notebook_principal.select(self.aba_sistema)
        tela.lift()
        try:
            tela.carregar_dados()
        except Exception:
            pass

    def voltar_tela(self):
        if self.historico_telas:
            nome = self.historico_telas.pop()
            self.mostrar_tela(nome, registrar=False)

    def sair_da_conta(self):
        self.usuario_logado = None
        self.nome_usuario_logado = ''
        self.historico_telas.clear()
        self.mostrar_tela(
            self.nome_tela_login, registrar=False
        )

    def fechar(self):
        try:
            self.conexao.close()
        finally:
            self.destroy()


class Tela1_TelaPrincipal(ttk.Frame):
    def __init__(self, master, app):
        super().__init__(master, style='App.TFrame')
        self.app = app
        self._id_selecionado = None
        self._criar_banco()
        self._criar_interface()
        self._atualizar_contadores()
        self.carregar_dados()


    def _criar_area_rolavel(self, largura, altura):
        self._largura_area = largura
        self._altura_area = altura
        self.canvas_rolavel = tk.Canvas(self, highlightthickness=0, bg=self.app.cor_fundo)
        self.barra_vertical = ttk.Scrollbar(self, orient='vertical', command=self.canvas_rolavel.yview)
        self.canvas_rolavel.configure(yscrollcommand=self.barra_vertical.set)
        self.canvas_rolavel.pack(side='left', fill='both', expand=True)
        self.barra_vertical.pack(side='right', fill='y')
        self.area = ttk.Frame(self.canvas_rolavel, style='App.TFrame', width=largura, height=altura)
        self.area.pack_propagate(False)
        self._janela_area = self.canvas_rolavel.create_window((0, 8), window=self.area, anchor='nw')
        self.canvas_rolavel.bind('<Configure>', self._ajustar_area_rolavel)
        self.area.bind('<Configure>', self._atualizar_scroll_area)
        self.canvas_rolavel.bind('<Enter>', lambda e: self.canvas_rolavel.bind_all('<MouseWheel>', self._ao_rolar_mouse))
        self.canvas_rolavel.bind('<Leave>', lambda e: self.canvas_rolavel.unbind_all('<MouseWheel>'))
        self.after(60, self._ajustar_area_rolavel)

    def _ajustar_area_rolavel(self, event=None):
        largura_canvas = max(1, self.canvas_rolavel.winfo_width())
        x = 0 if largura_canvas <= self._largura_area else int((largura_canvas - self._largura_area) / 2)
        self.canvas_rolavel.coords(self._janela_area, x, 8)
        self._atualizar_scroll_area()

    def _atualizar_scroll_area(self, event=None):
        largura_total = max(self._largura_area, self.canvas_rolavel.winfo_width())
        altura_total = max(self._altura_area + 16, self.canvas_rolavel.winfo_height())
        self.canvas_rolavel.configure(scrollregion=(0, 0, largura_total, altura_total))

    def _ao_rolar_mouse(self, event):
        try:
            deslocamento = int(-1 * (event.delta / 120))
            if deslocamento != 0:
                self.canvas_rolavel.yview_scroll(deslocamento, 'units')
        except Exception:
            pass

    def _criar_banco(self):
        self.app.conexao.execute('CREATE TABLE IF NOT EXISTS registros (id INTEGER PRIMARY KEY AUTOINCREMENT, nome TEXT, curso TEXT, idade TEXT, e_mail TEXT, sexo TEXT, telefone TEXT)')
        colunas_existentes = {linha[1] for linha in self.app.conexao.execute('PRAGMA table_info(registros)').fetchall()}
        if 'nome' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN nome TEXT')
        if 'curso' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN curso TEXT')
        if 'idade' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN idade TEXT')
        if 'e_mail' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN e_mail TEXT')
        if 'sexo' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN sexo TEXT')
        if 'telefone' not in colunas_existentes:
            self.app.conexao.execute('ALTER TABLE registros ADD COLUMN telefone TEXT')
        self.app.conexao.commit()

    def _criar_interface(self):
        self._criar_area_rolavel(1131, 748)
        self.label_1 = ttk.Label(self.area, text='Cadastro de Alunos', style=self.app._estilo_componente('Componenteitem_3caed07950.TLabel', 'TLabel', ('Segoe UI', 15, 'bold'), '', ''))
        self.label_1.place(x=12, y=0, width=273, height=30)
        self.label_2 = ttk.Label(self.area, text='preencha os campos abaixo', style=self.app._estilo_componente('Componenteitem_708dd0faa2.TLabel', 'TLabel', ('Arial', 12), '', ''))
        self.label_2.place(x=12, y=47, width=361, height=30)
        self.nomelabel = ttk.Label(self.area, text='Nome', style='TLabel')
        self.nomelabel.place(x=12, y=153, width=180, height=30)
        self.var_nome = tk.StringVar()
        self.nome = ttk.Entry(self.area, textvariable=self.var_nome)
        self.nome.place(x=65, y=153, width=220, height=34)
        self.cursolabel = ttk.Label(self.area, text='curso', style='TLabel')
        self.cursolabel.place(x=12, y=216, width=180, height=30)
        self.var_curso = tk.StringVar()
        self.curso = ttk.Entry(self.area, textvariable=self.var_curso)
        self.curso.place(x=65, y=212, width=220, height=34)
        self.idadelabel = ttk.Label(self.area, text='Idade', style='TLabel')
        self.idadelabel.place(x=350, y=151, width=180, height=30)
        self.var_idade = tk.StringVar()
        self.idade = ttk.Entry(self.area, textvariable=self.var_idade)
        self.idade.place(x=403, y=147, width=220, height=34)
        self.emaillabel = ttk.Label(self.area, text='E_mail', style='TLabel')
        self.emaillabel.place(x=350, y=206, width=180, height=30)
        self.var_e_mail = tk.StringVar()
        self.e_mail = ttk.Entry(self.area, textvariable=self.var_e_mail)
        self.e_mail.place(x=403, y=206, width=220, height=34)
        self.var_sexo = tk.StringVar()
        self.sexo = ttk.Combobox(self.area, textvariable=self.var_sexo, values=['-', 'Masculino', 'Feminino'], state='readonly')
        self.sexo.place(x=65, y=263, width=220, height=34)
        self.sexolabel = ttk.Label(self.area, text='Sexo', style='TLabel')
        self.sexolabel.place(x=12, y=263, width=35, height=30)
        self.telefonelabel = ttk.Label(self.area, text='Telefone', style='TLabel')
        self.telefonelabel.place(x=350, y=259, width=62, height=30)
        self.var_telefone = tk.StringVar()
        self.telefone = ttk.Entry(self.area, textvariable=self.var_telefone)
        self.telefone.place(x=412, y=259, width=220, height=34)
        self.buttoncadastrar = ttk.Button(self.area, text='Cadastrar', command=lambda: self.salvar_registro(['nome', 'idade', 'curso', 'e_mail', 'sexo', 'telefone']))
        self.buttoncadastrar.place(x=12, y=351, width=130, height=38)
        self.button_2 = ttk.Button(self.area, text='Atualizar', command=lambda: self.atualizar_registro(['nome', 'idade', 'curso', 'e_mail', 'sexo', 'telefone']))
        self.button_2.place(x=167, y=351, width=130, height=38)
        self.button_3 = ttk.Button(self.area, text='Excluir', command=self.excluir_registro)
        self.button_3.place(x=316, y=351, width=130, height=38)
        self.button_4 = ttk.Button(self.area, text='Limpar', command=lambda: self.limpar_campos(['nome', 'idade', 'curso', 'e_mail', 'sexo', 'telefone']))
        self.button_4.place(x=457, y=351, width=130, height=38)
        self.tabela_dados = ttk.Treeview(self.area, columns=['nome', 'curso', 'idade', 'e_mail', 'sexo', 'telefone'], show='headings')
        self.tabela_dados.heading('nome', text='Nome')
        self.tabela_dados.column('nome', width=140, anchor='w')
        self.tabela_dados.heading('curso', text='Curso')
        self.tabela_dados.column('curso', width=140, anchor='w')
        self.tabela_dados.heading('idade', text='Idade')
        self.tabela_dados.column('idade', width=140, anchor='w')
        self.tabela_dados.heading('e_mail', text='E Mail')
        self.tabela_dados.column('e_mail', width=140, anchor='w')
        self.tabela_dados.heading('sexo', text='Sexo')
        self.tabela_dados.column('sexo', width=140, anchor='w')
        self.tabela_dados.heading('telefone', text='Telefone')
        self.tabela_dados.column('telefone', width=140, anchor='w')
        self.tabela_dados.place(x=0, y=443, width=725, height=230)
        self.tabela_dados.bind('<<TreeviewSelect>>', self.preencher_formulario)
        self.var_filtro = tk.StringVar()
        self.filtro = ttk.Entry(self.area, textvariable=self.var_filtro)
        self.filtro.place(x=0, y=401, width=240, height=34)
        self.filtro.bind('<KeyRelease>', lambda e: self.aplicar_filtro())

    def _valor_componente(self, nome):
        variavel = getattr(self, f'var_{nome}', None)
        if variavel is not None:
            return variavel.get()
        widget = getattr(self, nome, None)
        if isinstance(widget, tk.Text):
            return widget.get('1.0', 'end-1c')
        if isinstance(widget, tk.Listbox):
            selecao = widget.curselection()
            if selecao:
                return widget.get(selecao[0])
            return '\n'.join(widget.get(0, tk.END))
        return ''

    def _definir_valor_componente(self, nome, valor):
        valor = '' if valor is None else valor
        variavel = getattr(self, f'var_{nome}', None)
        if variavel is not None:
            variavel.set(valor)
            return True
        widget = getattr(self, nome, None)
        if isinstance(widget, tk.Text):
            widget.delete('1.0', tk.END)
            texto = '\n'.join(str(x) for x in valor) if isinstance(valor, (list, tuple)) else str(valor)
            widget.insert('1.0', texto)
            return True
        if isinstance(widget, tk.Listbox):
            widget.delete(0, tk.END)
            itens = valor if isinstance(valor, (list, tuple)) else str(valor).splitlines()
            for item in itens:
                widget.insert(tk.END, item)
            return True
        return False

    def _atualizar_contadores(self):
        pass

    def exportar_dados_excel(self, alvos=None):
        try:
            from openpyxl import Workbook
        except ImportError:
            messagebox.showerror('Excel', 'Para exportar arquivos .xlsx, instale a biblioteca openpyxl com: pip install openpyxl')
            return
        nomes = list(alvos or [])
        if not nomes:
            nomes = [nome for nome, widget in self.__dict__.items() if isinstance(widget, (ttk.Treeview, tk.Listbox))]
        fontes = [(nome, getattr(self, nome, None)) for nome in nomes]
        fontes = [(nome, widget) for nome, widget in fontes if isinstance(widget, (ttk.Treeview, tk.Listbox))]
        if not fontes:
            messagebox.showwarning('Excel', 'Selecione uma Treeview ou uma Listbox nos Campos da ação.')
            return
        caminho = filedialog.asksaveasfilename(title='Exportar para Excel', defaultextension='.xlsx', filetypes=[('Arquivo do Excel', '*.xlsx')])
        if not caminho:
            return
        wb = Workbook()
        wb.remove(wb.active)
        def valor_seguro(valor):
            if isinstance(valor, str) and valor[:1] in ('=', '+', '-', '@'):
                return "'" + valor
            return valor
        for indice, (nome, widget) in enumerate(fontes, 1):
            titulo = re.sub(r'[:\\/?*\[\]]', '_', str(nome))[:31] or f'Dados{indice}'
            base = titulo
            contador = 2
            while titulo in wb.sheetnames:
                sufixo = f'_{contador}'
                titulo = base[:31-len(sufixo)] + sufixo
                contador += 1
            ws = wb.create_sheet(titulo)
            if isinstance(widget, ttk.Treeview):
                colunas = list(widget['columns'])
                cabecalhos = [widget.heading(coluna, 'text') or str(coluna) for coluna in colunas]
                ws.append(cabecalhos)
                for item in widget.get_children():
                    ws.append([valor_seguro(v) for v in widget.item(item, 'values')])
            else:
                ws.append(['Item'])
                for item in widget.get(0, tk.END):
                    ws.append([valor_seguro(item)])
            for coluna in ws.columns:
                letra = coluna[0].column_letter
                maior = max((len(str(celula.value or '')) for celula in coluna), default=10)
                ws.column_dimensions[letra].width = min(max(maior + 2, 12), 60)
        try:
            wb.save(caminho)
        except OSError as erro:
            messagebox.showerror('Excel', f'Não foi possível salvar o arquivo: {erro}')
            return
        messagebox.showinfo('Excel', 'Dados exportados para o Excel com sucesso!')

    def _formatar_tempo(self, segundos):
        segundos = max(0.0, float(segundos))
        horas = int(segundos // 3600)
        minutos = int((segundos % 3600) // 60)
        resto = segundos % 60
        return f'{horas:02d}:{minutos:02d}:{resto:04.1f}'

    def cronometro_iniciar(self, destino):
        if getattr(self, '_cronometro_rodando', False):
            return
        self._cronometro_rodando = True
        self._cronometro_acumulado = float(getattr(self, '_cronometro_acumulado', 0.0))
        self._cronometro_inicio = time.perf_counter()
        self._cronometro_tick(destino)

    def _cronometro_tick(self, destino):
        if not getattr(self, '_cronometro_rodando', False):
            return
        atual = self._cronometro_acumulado + (time.perf_counter() - self._cronometro_inicio)
        self._definir_valor_componente(destino, self._formatar_tempo(atual))
        self.after(100, lambda: self._cronometro_tick(destino))

    def cronometro_pausar(self, destino):
        if getattr(self, '_cronometro_rodando', False):
            self._cronometro_acumulado += time.perf_counter() - self._cronometro_inicio
        self._cronometro_rodando = False
        self._definir_valor_componente(destino, self._formatar_tempo(getattr(self, '_cronometro_acumulado', 0.0)))

    def cronometro_zerar(self, destino):
        self._cronometro_rodando = False
        self._cronometro_acumulado = 0.0
        self._definir_valor_componente(destino, '00:00:00.0')

    def temporizador_iniciar(self, origem, destino):
        if getattr(self, '_temporizador_rodando', False):
            return
        try:
            valor = int(float(str(self._valor_componente(origem)).replace(',', '.')))
        except (TypeError, ValueError):
            valor = 0
        if valor <= 0:
            messagebox.showwarning('Temporizador', 'Informe uma quantidade de segundos maior que zero.')
            return
        self._temporizador_restante = valor
        self._temporizador_rodando = True
        self._temporizador_tick(destino)

    def _temporizador_tick(self, destino):
        if not getattr(self, '_temporizador_rodando', False):
            return
        restante = int(getattr(self, '_temporizador_restante', 0))
        self._definir_valor_componente(destino, self._formatar_tempo(restante))
        if restante <= 0:
            self._temporizador_rodando = False
            messagebox.showinfo('Temporizador', 'Tempo encerrado!')
            return
        self._temporizador_restante = restante - 1
        self.after(1000, lambda: self._temporizador_tick(destino))

    def temporizador_pausar(self, destino):
        self._temporizador_rodando = False
        self._definir_valor_componente(destino, self._formatar_tempo(getattr(self, '_temporizador_restante', 0)))

    def temporizador_zerar(self, destino):
        self._temporizador_rodando = False
        self._temporizador_restante = 0
        self._definir_valor_componente(destino, '00:00:00.0')

    def gerar_senha_segura(self, campo_tamanho, campo_simbolos, destino):
        try:
            tamanho = int(float(str(self._valor_componente(campo_tamanho)).replace(',', '.')))
        except (TypeError, ValueError):
            tamanho = 12
        tamanho = min(max(tamanho, 8), 128)
        usar_simbolos = bool(self._valor_componente(campo_simbolos)) if campo_simbolos else True
        alfabeto = string.ascii_letters + string.digits
        if usar_simbolos:
            alfabeto += '!@#$%&*_-+=?'
        senha = ''.join(secrets.choice(alfabeto) for _ in range(tamanho))
        self._definir_valor_componente(destino, senha)

    def calcular_idade_completa(self, origem, destino):
        texto = str(self._valor_componente(origem)).strip()
        try:
            nascimento = datetime.strptime(texto, '%d/%m/%Y').date()
        except ValueError:
            messagebox.showwarning('Idade', 'Informe a data de nascimento no formato dd/mm/aaaa.')
            return
        hoje = date.today()
        if nascimento > hoje:
            messagebox.showwarning('Idade', 'A data de nascimento não pode estar no futuro.')
            return
        anos = hoje.year - nascimento.year
        meses = hoje.month - nascimento.month
        dias = hoje.day - nascimento.day
        if dias < 0:
            meses -= 1
            mes_anterior = hoje.month - 1 or 12
            ano_mes_anterior = hoje.year if hoje.month > 1 else hoje.year - 1
            dias += calendar.monthrange(ano_mes_anterior, mes_anterior)[1]
        if meses < 0:
            anos -= 1
            meses += 12
        self._definir_valor_componente(destino, f'{anos} anos, {meses} meses e {dias} dias')

    def sortear_nome(self, origem, destino):
        widget = getattr(self, origem, None)
        if isinstance(widget, tk.Listbox):
            nomes = [str(x).strip() for x in widget.get(0, tk.END) if str(x).strip()]
        else:
            nomes = [x.strip() for x in str(self._valor_componente(origem)).splitlines() if x.strip()]
        if not nomes:
            messagebox.showwarning('Sorteio', 'A lista de nomes está vazia.')
            return
        self._definir_valor_componente(destino, random.choice(nomes))

    def gerar_numeros_sorteio(self, origem, destino):
        try:
            quantidade = int(float(str(self._valor_componente(origem)).replace(',', '.')))
        except (TypeError, ValueError):
            quantidade = 6
        if quantidade <= 0:
            quantidade = 6
        quantidade = min(quantidade, 60)
        numeros = sorted(random.sample(range(1, 61), quantidade))
        self._definir_valor_componente(destino, [f'{numero:02d}' for numero in numeros])
        self._atualizar_contadores()

    def validar_formulario(self):
        return True

    def _aplicar_mascara(self, nome, mascara):
        valor = str(self._valor_componente(nome))
        digitos = re.sub(r'\D', '', valor)
        novo = valor
        if mascara == 'CPF' and len(digitos) >= 11:
            d = digitos[:11]; novo = f'{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}'
        elif mascara == 'CNPJ' and len(digitos) >= 14:
            d = digitos[:14]; novo = f'{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}'
        elif mascara == 'Telefone' and len(digitos) >= 10:
            d = digitos[:11]; novo = f'({d[:2]}) {d[2:7]}-{d[7:]}' if len(d) == 11 else f'({d[:2]}) {d[2:6]}-{d[6:]}'
        elif mascara == 'CEP' and len(digitos) >= 8:
            d = digitos[:8]; novo = f'{d[:5]}-{d[5:]}'
        elif mascara == 'Data' and len(digitos) >= 8:
            d = digitos[:8]; novo = f'{d[:2]}/{d[2:4]}/{d[4:]}'
        elif mascara == 'Hora' and len(digitos) >= 4:
            d = digitos[:4]; novo = f'{d[:2]}:{d[2:4]}'
        elif mascara == 'Moeda':
            try:
                numero = float(valor.replace('R$', '').replace('.', '').replace(',', '.').strip())
                novo = f'R$ {numero:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
            except ValueError:
                return
        self._definir_valor_componente(nome, novo)

    def validar_login(self):
        messagebox.showinfo('Login', 'Esta tela não está configurada como Login.')

    def salvar_registro(self, selecionados=None):
        if not self.validar_formulario():
            return
        mapa = {'nome': 'nome', 'curso': 'curso', 'idade': 'idade', 'e_mail': 'e_mail', 'sexo': 'sexo', 'telefone': 'telefone'}
        escolhidos = list(selecionados or ['nome', 'curso', 'idade', 'e_mail', 'sexo', 'telefone'])
        escolhidos = [nome for nome in escolhidos if nome in mapa]
        if not escolhidos:
            messagebox.showwarning('SQLite', 'Selecione pelo menos um campo para cadastrar.')
            return
        campos = [mapa[nome] for nome in escolhidos]
        valores = []
        for nome, campo in zip(escolhidos, campos):
            valor = self._valor_componente(nome)
            if campo == 'senha' and str(valor):
                valor = hashlib.sha256(str(valor).encode('utf-8')).hexdigest()
            valores.append(valor)
        placeholders = ','.join('?' for _ in campos)
        sql = 'INSERT INTO registros (' + ','.join(campos) + ') VALUES (' + placeholders + ')' 
        try:
            self.app.conexao.execute(sql, valores)
            self.app.conexao.commit()
        except sqlite3.IntegrityError as erro:
            messagebox.showerror('Banco de dados', f'Não foi possível cadastrar: {erro}')
            return
        self.carregar_dados()
        messagebox.showinfo('SQLite', 'Dados cadastrados com sucesso!')

    def carregar_dados(self):
        tabela = self.tabela_dados
        for item in tabela.get_children():
            tabela.delete(item)
        linhas = self.app.conexao.execute('SELECT id,nome,curso,idade,e_mail,sexo,telefone FROM registros ORDER BY id DESC').fetchall()
        for linha in linhas:
            tabela.insert('', tk.END, iid=str(linha[0]), values=linha[1:])
        self._atualizar_contadores()

    def aplicar_filtro(self):
        termo = self.var_filtro.get().strip()
        tabela = self.tabela_dados
        for item in tabela.get_children():
            tabela.delete(item)
        if termo:
            sql = 'SELECT id,nome,curso,idade,e_mail,sexo,telefone FROM registros WHERE CAST(nome AS TEXT) LIKE ? ORDER BY id DESC'
            linhas = self.app.conexao.execute(sql, (f'%{termo}%',)).fetchall()
        else:
            linhas = self.app.conexao.execute('SELECT id,nome,curso,idade,e_mail,sexo,telefone FROM registros ORDER BY id DESC').fetchall()
        for linha in linhas:
            tabela.insert('', tk.END, iid=str(linha[0]), values=linha[1:])
        self._atualizar_contadores()

    def preencher_formulario(self, event=None):
        selecionado = self.tabela_dados.selection()
        if not selecionado:
            return
        valores = self.tabela_dados.item(selecionado[0], 'values')
        if not valores:
            return
        self._id_selecionado = selecionado[0]
        if len(valores) > 0:
            self._definir_valor_componente('nome', valores[0])
        if len(valores) > 1:
            self._definir_valor_componente('curso', valores[1])
        if len(valores) > 2:
            self._definir_valor_componente('idade', valores[2])
        if len(valores) > 3:
            self._definir_valor_componente('e_mail', valores[3])
        if len(valores) > 4:
            self._definir_valor_componente('sexo', valores[4])
        if len(valores) > 5:
            self._definir_valor_componente('telefone', valores[5])

    def atualizar_registro(self, selecionados=None):
        if self._id_selecionado is None:
            messagebox.showwarning('Atenção', 'Selecione um registro na tabela da tela antes de atualizar.')
            return
        if not self.validar_formulario():
            return
        mapa = {'nome': 'nome', 'curso': 'curso', 'idade': 'idade', 'e_mail': 'e_mail', 'sexo': 'sexo', 'telefone': 'telefone'}
        escolhidos = list(selecionados or ['nome', 'curso', 'idade', 'e_mail', 'sexo', 'telefone'])
        escolhidos = [nome for nome in escolhidos if nome in mapa]
        if not escolhidos:
            messagebox.showwarning('SQLite', 'Selecione os campos que serão atualizados.')
            return
        campos = [mapa[nome] for nome in escolhidos]
        pares_atualizacao = []
        for nome, campo in zip(escolhidos, campos):
            valor = self._valor_componente(nome)
            if campo == 'senha':
                if not str(valor):
                    continue
                valor = hashlib.sha256(str(valor).encode('utf-8')).hexdigest()
            pares_atualizacao.append((campo, valor))
        if not pares_atualizacao:
            messagebox.showwarning('SQLite', 'Nenhum campo possui um novo valor para atualizar.')
            return
        campos = [campo for campo, _ in pares_atualizacao]
        valores = [valor for _, valor in pares_atualizacao]
        sets = ', '.join(f'{campo} = ?' for campo in campos)
        self.app.conexao.execute('UPDATE registros SET ' + sets + ' WHERE id = ?', valores + [self._id_selecionado])
        self.app.conexao.commit()
        self.carregar_dados()
        messagebox.showinfo('SQLite', 'Registro atualizado com sucesso!')

    def excluir_registro(self):
        selecionado = self.tabela_dados.selection()
        if not selecionado:
            messagebox.showwarning('Atenção', 'Selecione um registro na tabela.')
            return
        valores = self.tabela_dados.item(selecionado[0], 'values')
        if not valores:
            return
        if not messagebox.askyesno('Excluir', 'Deseja realmente excluir o registro selecionado?'):
            return
        self.app.conexao.execute('DELETE FROM registros WHERE id = ?', (selecionado[0],))
        self.app.conexao.commit()
        self.carregar_dados()
        self.limpar_campos()

    def limpar_campos(self, selecionados=None):
        self._id_selecionado = None
        nomes = list(selecionados or ['nome', 'curso', 'idade', 'e_mail', 'sexo', 'telefone'])
        for nome in nomes:
            variavel = getattr(self, f'var_{nome}', None)
            if variavel is not None:
                try:
                    variavel.set(False if isinstance(variavel, tk.BooleanVar) else '')
                except tk.TclError:
                    pass
                continue
            widget = getattr(self, nome, None)
            if isinstance(widget, tk.Text):
                widget.delete('1.0', tk.END)
            elif isinstance(widget, tk.Listbox):
                widget.delete(0, tk.END)

    def _valor_para_formula(self, valor):
        if isinstance(valor, (int, float, bool)):
            return valor
        original = '' if valor is None else str(valor).strip()
        texto = original.replace('R$', '').replace(' ', '')
        try:
            if ',' in texto:
                texto = texto.replace('.', '').replace(',', '.')
            numero = float(texto)
            return int(numero) if numero.is_integer() else numero
        except (ValueError, TypeError):
            return original

    def _formatar_resultado(self, valor):
        if isinstance(valor, float):
            if valor.is_integer():
                return str(int(valor))
            return f'{valor:.10f}'.rstrip('0').rstrip('.').replace('.', ',')
        return str(valor)

    def _avaliar_expressao_segura(self, expressao, valores=None, resultado_anterior=None, n=None):
        valores = list(valores or [])
        def nome_campo_formula(indice):
            numero = indice + 1
            partes = []
            while numero > 0:
                numero, resto = divmod(numero - 1, 26)
                partes.append(chr(65 + resto))
            return ''.join(reversed(partes))
        ambiente = {nome_campo_formula(i): valor for i, valor in enumerate(valores)}
        ambiente.update({f'CAMPO{i + 1}': valor for i, valor in enumerate(valores)})
        if resultado_anterior is not None: ambiente['R'] = resultado_anterior
        if n is not None: ambiente['N'] = n
        funcoes = {
            'round': round, 'abs': abs, 'min': min, 'max': max,
            'int': int, 'float': float, 'str': str,
            'upper': lambda x: str(x).upper(),
            'lower': lambda x: str(x).lower(), 'len': len
        }
        arvore = ast.parse(str(expressao), mode='eval')

        def visitar(no):
            if isinstance(no, ast.Expression):
                return visitar(no.body)
            if isinstance(no, ast.Constant):
                return no.value
            if isinstance(no, ast.Name):
                if no.id in ambiente:
                    return ambiente[no.id]
                raise ValueError(f'Nome não permitido: {no.id}')
            if isinstance(no, ast.BinOp):
                a, b = visitar(no.left), visitar(no.right)
                if isinstance(no.op, ast.Add): return a + b
                if isinstance(no.op, ast.Sub): return a - b
                if isinstance(no.op, ast.Mult): return a * b
                if isinstance(no.op, ast.Div): return a / b
                if isinstance(no.op, ast.FloorDiv): return a // b
                if isinstance(no.op, ast.Mod): return a % b
                if isinstance(no.op, ast.Pow): return a ** b
                raise ValueError('Operador não permitido')
            if isinstance(no, ast.UnaryOp):
                valor = visitar(no.operand)
                if isinstance(no.op, ast.USub): return -valor
                if isinstance(no.op, ast.UAdd): return +valor
                if isinstance(no.op, ast.Not): return not valor
                raise ValueError('Operador unário não permitido')
            if isinstance(no, ast.BoolOp):
                valores_bool = [visitar(x) for x in no.values]
                return all(valores_bool) if isinstance(no.op, ast.And) else any(valores_bool)
            if isinstance(no, ast.Compare):
                esquerda = visitar(no.left)
                for operador, comparador in zip(no.ops, no.comparators):
                    direita = visitar(comparador)
                    if isinstance(operador, ast.Eq): ok = esquerda == direita
                    elif isinstance(operador, ast.NotEq): ok = esquerda != direita
                    elif isinstance(operador, ast.Gt): ok = esquerda > direita
                    elif isinstance(operador, ast.GtE): ok = esquerda >= direita
                    elif isinstance(operador, ast.Lt): ok = esquerda < direita
                    elif isinstance(operador, ast.LtE): ok = esquerda <= direita
                    else: raise ValueError('Comparação não permitida')
                    if not ok: return False
                    esquerda = direita
                return True
            if isinstance(no, ast.IfExp):
                return visitar(no.body) if visitar(no.test) else visitar(no.orelse)
            if isinstance(no, ast.Call):
                if not isinstance(no.func, ast.Name) or no.func.id not in funcoes:
                    raise ValueError('Função não permitida')
                if no.keywords:
                    raise ValueError('Argumentos nomeados não são permitidos')
                return funcoes[no.func.id](*[visitar(x) for x in no.args])
            raise ValueError(f'Elemento não permitido: {type(no).__name__}')

        return visitar(arvore)

    def _operacao_basica(self, valores, operador, expressao=''):
        if operador == 'Expressão personalizada':
            if not str(expressao).strip():
                raise ValueError('Informe a expressão personalizada.')
            return self._avaliar_expressao_segura(expressao, valores)
        if operador == 'Operação definida no Campo 3':
            if len(valores) < 3:
                raise ValueError('Selecione Campo 1, Campo 2 e Campo 3.')
            op = str(valores[2]).strip().lower()
            mapa = {
                'somar': 'Somar', '+': 'Somar',
                'subtrair': 'Subtrair', '-': 'Subtrair',
                'multiplicar': 'Multiplicar', '*': 'Multiplicar', '×': 'Multiplicar',
                'dividir': 'Dividir', '/': 'Dividir', '÷': 'Dividir'
            }
            operador = mapa.get(op, '')
            valores = valores[:2]
        if not valores:
            raise ValueError('Selecione pelo menos um campo.')
        if operador == 'Somar':
            return sum(valores)
        if operador == 'Subtrair':
            resultado = valores[0]
            for valor in valores[1:]: resultado -= valor
            return resultado
        if operador == 'Multiplicar':
            resultado = 1
            for valor in valores: resultado *= valor
            return resultado
        if operador == 'Dividir':
            resultado = valores[0]
            for valor in valores[1:]: resultado /= valor
            return resultado
        if operador == 'Média':
            return sum(valores) / len(valores)
        if operador == 'Porcentagem':
            if len(valores) < 2: raise ValueError('Porcentagem precisa de dois campos.')
            return valores[0] * valores[1] / 100
        if operador == 'Resto da divisão':
            if len(valores) < 2: raise ValueError('Resto precisa de dois campos.')
            return valores[0] % valores[1]
        if operador == 'Potência':
            if len(valores) < 2: raise ValueError('Potência precisa de dois campos.')
            return valores[0] ** valores[1]
        if operador == 'Mínimo':
            return min(valores)
        if operador == 'Máximo':
            return max(valores)
        raise ValueError('Escolha uma operação matemática.')

    def inserir_valor(self, destino, valor, modo='Adicionar'):
        if not destino:
            messagebox.showwarning('Ação', 'Selecione o campo que receberá o valor.')
            return
        atual = str(self._valor_componente(destino))
        novo = str(valor) if modo == 'Substituir' else atual + str(valor)
        if not self._definir_valor_componente(destino, novo):
            messagebox.showwarning('Ação', 'O campo de destino não existe nesta tela.')

    def calcular_campos(self, campos, destino, operador, expressao='', formato='', destino_2='', expressao_2=''):
        if not campos or not destino:
            messagebox.showwarning('Cálculo', 'Selecione os campos da operação e o campo de resultado.')
            return
        try:
            valores = [self._valor_para_formula(self._valor_componente(nome)) for nome in campos]
            resultado = self._operacao_basica(valores, operador, expressao)
            texto = self._formatar_resultado(resultado)
            if formato:
                texto = str(formato).replace('{resultado}', texto)
            if not self._definir_valor_componente(destino, texto):
                raise ValueError('Campo de resultado não encontrado.')
            if destino_2 and expressao_2:
                resultado_2 = self._avaliar_expressao_segura(expressao_2, valores, resultado_anterior=resultado)
                self._definir_valor_componente(destino_2, self._formatar_resultado(resultado_2))
        except ZeroDivisionError:
            messagebox.showwarning('Cálculo', 'Não é possível dividir por zero.')
        except Exception as erro:
            messagebox.showwarning('Cálculo', f'Não foi possível executar a operação: {erro}')

    def calcular_expressao_campo(self, origem, destino, auxiliar=''):
        if not origem or not destino:
            messagebox.showwarning('Cálculo', 'Selecione o campo da expressão e o campo de resultado.')
            return
        original = str(self._valor_componente(origem)).strip()
        expressao = original.replace('×', '*').replace('÷', '/').replace(',', '.')
        if not expressao:
            return
        try:
            resultado = self._avaliar_expressao_segura(expressao)
            if not isinstance(resultado, (int, float)):
                raise ValueError('A expressão deve produzir um número.')
            texto = self._formatar_resultado(resultado)
            self._definir_valor_componente(destino, texto)
            if auxiliar:
                widget = getattr(self, auxiliar, None)
                if isinstance(widget, tk.Listbox):
                    widget.insert(tk.END, f'{original} = {texto}')
                elif isinstance(widget, tk.Text):
                    widget.insert(tk.END, f'{original} = {texto}\n')
        except ZeroDivisionError:
            messagebox.showwarning('Cálculo', 'Não é possível dividir por zero.')
        except Exception as erro:
            messagebox.showwarning('Cálculo', f'Expressão inválida: {erro}')

    def aplicar_condicao(self, campos, destino, comparador, valor_fixo='', verdadeiro='Sim', falso='Não'):
        if not campos or not destino or not comparador:
            messagebox.showwarning('Condição', 'Selecione Campo 1, comparador e campo de resultado.')
            return
        a = self._valor_para_formula(self._valor_componente(campos[0]))
        if len(campos) >= 2:
            b = self._valor_para_formula(self._valor_componente(campos[1]))
        else:
            b = self._valor_para_formula(valor_fixo)
        try:
            if comparador == '==': ok = a == b
            elif comparador == '!=': ok = a != b
            elif comparador == '>': ok = a > b
            elif comparador == '>=': ok = a >= b
            elif comparador == '<': ok = a < b
            elif comparador == '<=': ok = a <= b
            else: raise ValueError('Comparador inválido')
            self._definir_valor_componente(destino, verdadeiro if ok else falso)
        except Exception as erro:
            messagebox.showwarning('Condição', f'Não foi possível comparar os valores: {erro}')

    def gerar_lista_repeticao(self, campos, destino, operador, inicio='1', fim='10', modelo=''):
        if not campos or not destino:
            messagebox.showwarning('Repetição', 'Selecione Campo 1 e o campo de resultado.')
            return
        try:
            a = self._valor_para_formula(self._valor_componente(campos[0]))
            inicio_i, fim_i = int(inicio), int(fim)
            passo = 1 if fim_i >= inicio_i else -1
            linhas_saida = []
            for n in range(inicio_i, fim_i + passo, passo):
                resultado = self._operacao_basica([a, n], operador, '')
                texto_resultado = self._formatar_resultado(resultado)
                linha = modelo or '{A} {operador} {N} = {resultado}'
                linha = str(linha).replace('{A}', self._formatar_resultado(a)).replace('{N}', str(n))
                linha = linha.replace('{operador}', operador).replace('{resultado}', texto_resultado)
                linhas_saida.append(linha)
            self._definir_valor_componente(destino, linhas_saida)
        except Exception as erro:
            messagebox.showwarning('Repetição', f'Não foi possível gerar a lista: {erro}')



if __name__ == '__main__':
    app = AplicacaoGerada()
    app.protocol('WM_DELETE_WINDOW', app.fechar)
    app.mainloop()
